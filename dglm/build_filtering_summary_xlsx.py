#!/usr/bin/env python3
"""
build_filtering_summary_xlsx.py

Builds a real Excel workbook summarizing the condition-level filtering
(cells/animals per subcluster x region, kept/removed at 300 and 100 cell
thresholds) for all 12 cell types. Reads directly from the already-computed
sparsity CSVs -- no new analysis, just compiling existing real data.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

CELL_TYPES = ['astrocytes', 'basket_cells', 'cerebellar_neurons', 'ependymal_cells',
              'GABAergic_neurons', 'glutamatergic_neurons', 'medium_spiny_neurons',
              'microglia', 'midbrain_neurons', 'opc', 'oligodendrocytes', 'vascular_cells']

BASE = '/scratch/easmit31/dispersion/dglm/disp_age__condition_sparsity_diagnostics'
OUT = '/scratch/easmit31/dispersion/dglm/filtering_summary_all_celltypes.xlsx'

FONT = Font(name='Arial', size=10)
HEADER_FONT = Font(name='Arial', size=10, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')

wb = Workbook()
summary_ws = wb.active
summary_ws.title = 'Summary'

detail_sheet_names = {}
for ct in CELL_TYPES:
    d300 = pd.read_csv(f'{BASE}/min300cells/{ct}_condition_sparsity.csv')
    d100 = pd.read_csv(f'{BASE}/min100cells/{ct}_condition_sparsity.csv')
    d300 = d300.rename(columns={'pass_both': 'pass_300c'})[['subcluster', 'region', 'n_cells', 'n_animals', 'pass_300c']]
    d100 = d100.rename(columns={'pass_both': 'pass_100c'})[['subcluster', 'region', 'pass_100c']]
    merged = d300.merge(d100, on=['subcluster', 'region'], how='outer')

    sheet_name = ct[:31]
    ws = wb.create_sheet(sheet_name)

    headers = ['subcluster', 'region', 'n_cells', 'n_animals', 'pass_300c', 'pass_100c']
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal='center')

    for r, row in enumerate(merged.itertuples(index=False), 2):
        ws.cell(row=r, column=1, value=row.subcluster).font = FONT
        ws.cell(row=r, column=2, value=row.region).font = FONT
        ws.cell(row=r, column=3, value=int(row.n_cells)).font = FONT
        ws.cell(row=r, column=4, value=int(row.n_animals)).font = FONT
        ws.cell(row=r, column=5, value=bool(row.pass_300c)).font = FONT
        ws.cell(row=r, column=6, value=bool(row.pass_100c)).font = FONT

    for col, width in zip('ABCDEF', [22, 10, 10, 10, 11, 11]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A2'

    detail_sheet_names[ct] = (sheet_name, len(merged) + 1)

headers = ['Cell type', 'Total conditions', 'Kept @300c', 'Removed @300c', '% removed @300c',
           'Kept @100c', 'Removed @100c', '% removed @100c',
           'Total cells', 'Cells kept @300c', 'Cells removed @300c', '% cells removed @300c',
           'Cells kept @100c', 'Cells removed @100c', '% cells removed @100c']
for col, h in enumerate(headers, 1):
    c = summary_ws.cell(row=1, column=col, value=h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal='center', wrap_text=True)
summary_ws.row_dimensions[1].height = 30

for i, ct in enumerate(CELL_TYPES):
    r = i + 2
    sheet_name, last_row = detail_sheet_names[ct]
    rng = f"'{sheet_name}'"

    summary_ws.cell(row=r, column=1, value=ct).font = FONT
    summary_ws.cell(row=r, column=2, value=f'=COUNTA({rng}!A2:A{last_row})').font = FONT
    summary_ws.cell(row=r, column=3, value=f'=COUNTIF({rng}!E2:E{last_row},TRUE)').font = FONT
    summary_ws.cell(row=r, column=4, value=f'=B{r}-C{r}').font = FONT
    summary_ws.cell(row=r, column=5, value=f'=D{r}/B{r}').font = FONT
    summary_ws.cell(row=r, column=5).number_format = '0.0%'
    summary_ws.cell(row=r, column=6, value=f'=COUNTIF({rng}!F2:F{last_row},TRUE)').font = FONT
    summary_ws.cell(row=r, column=7, value=f'=B{r}-F{r}').font = FONT
    summary_ws.cell(row=r, column=8, value=f'=G{r}/B{r}').font = FONT
    summary_ws.cell(row=r, column=8).number_format = '0.0%'
    summary_ws.cell(row=r, column=9, value=f'=SUM({rng}!C2:C{last_row})').font = FONT
    summary_ws.cell(row=r, column=10, value=f'=SUMIF({rng}!E2:E{last_row},TRUE,{rng}!C2:C{last_row})').font = FONT
    summary_ws.cell(row=r, column=11, value=f'=I{r}-J{r}').font = FONT
    summary_ws.cell(row=r, column=12, value=f'=K{r}/I{r}').font = FONT
    summary_ws.cell(row=r, column=12).number_format = '0.0%'
    summary_ws.cell(row=r, column=13, value=f'=SUMIF({rng}!F2:F{last_row},TRUE,{rng}!C2:C{last_row})').font = FONT
    summary_ws.cell(row=r, column=14, value=f'=I{r}-M{r}').font = FONT
    summary_ws.cell(row=r, column=15, value=f'=N{r}/I{r}').font = FONT
    summary_ws.cell(row=r, column=15).number_format = '0.0%'

for col, width in zip('ABCDEFGHIJKLMNO', [22,14,10,12,12,10,12,12,11,13,14,14,13,14,14]):
    summary_ws.column_dimensions[col].width = width
summary_ws.freeze_panes = 'A2'

wb.save(OUT)
print(f'Saved: {OUT}')
